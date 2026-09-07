from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path

OUT = Path(r"D:\Unity project\PhysicscollGame\Unity游戏开发秋招简历.docx")

BLUE = RGBColor(31, 78, 121)
INK = RGBColor(30, 37, 45)
MUTED = RGBColor(92, 101, 112)
LIGHT = "D9E4F0"
FONT_CN = "Microsoft YaHei"
FONT_LATIN = "Arial"


def set_run(run, size=9.4, bold=False, color=INK, italic=False):
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=40, start=80, bottom=40, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_bottom_rule(paragraph, color="B7C9DC", size="10"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def create_bullet_num(doc):
    numbering = doc.part.numbering_part.element
    abstract_id = 42
    num_id = 42
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl.append(lvl_text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "360")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "360")
    ind.set(qn("w:hanging"), "180")
    p_pr.append(ind)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_LATIN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_bullet(doc, num_id, parts):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.8)
    p.paragraph_format.line_spacing = 1.04
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)
    for text, bold in parts:
        set_run(p.add_run(text), 9.45, bold=bold)
    return p


def add_section(doc, title):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.keep_with_next = True
    p.add_run(title)
    return p


def add_project_header(doc, title, role, tech):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1.2)
    p.paragraph_format.space_after = Pt(1.0)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(title), 10.2, bold=True, color=INK)
    set_run(p.add_run(f"  |  {role}"), 9.2, bold=True, color=BLUE)
    set_run(p.add_run(f"\n{tech}"), 8.5, color=MUTED)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.48)
section.bottom_margin = Inches(0.45)
section.left_margin = Inches(0.62)
section.right_margin = Inches(0.62)
section.header_distance = Inches(0.2)
section.footer_distance = Inches(0.2)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT_LATIN
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
normal.font.size = Pt(9.7)
normal.font.color.rgb = INK
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(1.8)
normal.paragraph_format.line_spacing = 1.04

h1 = styles["Heading 1"]
h1.font.name = FONT_LATIN
h1._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
h1.font.size = Pt(11.8)
h1.font.bold = True
h1.font.color.rgb = BLUE
h1.paragraph_format.space_before = Pt(4.0)
h1.paragraph_format.space_after = Pt(2.0)
h1.paragraph_format.keep_with_next = True

num_id = create_bullet_num(doc)

# Header block: deliberately plain and editable because personal facts were not supplied.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
set_run(p.add_run("姓名（待补充）"), 18.5, bold=True, color=INK)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
set_run(p.add_run("Unity 游戏客户端开发 / Unity 游戏开发工程师"), 10.5, bold=True, color=BLUE)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2.5)
set_run(p.add_run("手机（待补充）  |  邮箱（待补充）  |  所在城市（待补充）  |  GitHub/作品链接（待补充）"), 8.6, color=MUTED)
add_bottom_rule(p)

add_section(doc, "教育背景")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
set_run(p.add_run("学校 / 专业 / 学历 / 毕业时间（待补充）"), 9.2, bold=True)

add_section(doc, "技术能力")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
set_run(p.add_run("Unity / C#："), 9.4, bold=True, color=BLUE)
set_run(p.add_run("Unity 2022 LTS、MonoBehaviour 生命周期、UGUI/TMP、Transform、Rigidbody2D/Collider2D、Animator 基础；能将游戏规则与表现层拆分。"), 9.4)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
set_run(p.add_run("游戏系统与网络："), 9.4, bold=True, color=BLUE)
set_run(p.add_run("固定 tick 模拟、自定义 2D 碰撞、状态机、对象复用、UDP 输入同步与有限窗口回滚原型、EditMode/PlayMode 测试。"), 9.4)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
set_run(p.add_run("DOTS / 算法："), 9.4, bold=True, color=BLUE)
set_run(p.add_run("了解 Entities 1.x、Baker、ISystem、DynamicBuffer、Burst/Job、NativeContainer；能分析 A*、最小堆、KD-Tree KNN 与 ORCA 数据流及性能边界。"), 9.4)

add_section(doc, "项目经历")
add_project_header(
    doc,
    "Pocket Bump｜2D 双人物理竞技游戏",
    "独立 Unity 开发",
    "Unity 2022.3 / C# / UGUI + TMP / UDP / 自定义固定步长模拟 / Unity Test Framework",
)
add_bullet(doc, num_id, [
    ("项目背景：", True),
    ("围绕双人争抢弹珠、篮筐蓄力射击与场地得分构建同机/局域网对战；将比赛推进、碰撞和计分从 Unity 组件中抽离为可复制的纯状态模拟。", False),
])
add_bullet(doc, num_id, [
    ("固定 tick 与碰撞：", True),
    ("以 60 Hz 推进 5 状态比赛流程，使用圆-圆、圆-AABB/OBB 和速度相关子步处理玩家/弹珠/篮筐碰撞；在 LateUpdate 插值前后状态，降低固定帧视觉抖动。", False),
])
add_bullet(doc, num_id, [
    ("回滚与协议：", True),
    ("针对 UDP 远端输入迟到，设计 2 tick 输入延迟、12 tick 快照环和上一输入预测；检测历史输入差异后恢复并重演，以 1 byte 压缩操作，并通过版本/规则握手、ack 与状态校验和发现连接异常和失同步。", False),
])
add_bullet(doc, num_id, [
    ("工程化与验证：", True),
    ("复用 1–8 个弹珠对象，避免回合内重复创建/销毁；按 tick 清理回滚后的瞬时特效；编写 29 个 EditMode/PlayMode 用例，覆盖回滚输入带、握手、碰撞/抢夺/计分和菜单流程（本次未在本机重新执行）。", False),
])

add_project_header(
    doc,
    "AStarEcs｜DOTS A* 源码研读与工程整理",
    "基于 Athomield 既有工程",
    "Entities 1.3 / Burst / Job System / Unity Physics / A* / KD-Tree KNN / ORCA",
)
add_bullet(doc, num_id, [
    ("架构梳理：", True),
    ("围绕 Component/Baker/System 整理网格烘焙、障碍与惩罚扫描、连通岛、请求队列、A* 路径回写和 LocalTransform 跟随的数据流；默认示例网格为 30×30（900 节点），每帧最多处理 25 个请求。", False),
])
add_bullet(doc, num_id, [
    ("算法与性能分析：", True),
    ("阅读 NativeMinHeap + NativeHashSet 的 A*、KD-Tree 批量 KNN 与 ORCA 速度约束；分析 330 Entity 圆阵示例，并识别同步 Complete、每帧 Native 临时分配及 open set 线性扫描等优化空间。KNN 为第三方 MIT 来源，不作为个人原创实现。", False),
])

# Quiet footer, no page number needed for a one-page resume.
footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p.paragraph_format.space_before = Pt(0)
set_run(footer_p.add_run("技术内容依据当前项目源码整理；投递前请补齐个人信息、项目时间与作品链接。"), 7.5, color=MUTED)

doc.core_properties.title = "Unity 游戏开发秋招简历"
doc.core_properties.subject = "Unity 游戏客户端开发 / Unity 游戏开发工程师"
doc.core_properties.author = ""
doc.core_properties.keywords = "Unity, C#, 游戏开发, 回滚, DOTS, ECS"

doc.save(OUT)
print(str(OUT))
